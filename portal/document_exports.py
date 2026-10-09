from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
from decimal import Decimal, InvalidOperation
import re

from django.conf import settings
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image as PdfImage
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from .models import _ad_soyad_bicimi
from .evrak_verisi import belge_verileri


STATIC_PORTAL = Path(__file__).resolve().parent.parent / "static" / "portal"
BELEDIYE_LOGO = STATIC_PORTAL / "trabzon-buyuksehir-logo.jpg"
TISKI_LOGO = STATIC_PORTAL / "logo1.png"
FONT_ADAYLARI = [
    (
        Path(getattr(settings, "PORTAL_PDF_FONT_REGULAR", "")),
        Path(getattr(settings, "PORTAL_PDF_FONT_BOLD", "")),
    ),
    (Path(r"C:\Windows\Fonts\arial.ttf"), Path(r"C:\Windows\Fonts\arialbd.ttf")),
    (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")),
    (Path("/Library/Fonts/Arial.ttf"), Path("/Library/Fonts/Arial Bold.ttf")),
]


def _fontlar():
    """Türkçe karakterleri PDF'de eksiksiz gösterecek yazı tiplerini kaydeder."""
    regular, bold = "Helvetica", "Helvetica-Bold"
    for regular_path, bold_path in FONT_ADAYLARI:
        if not regular_path or not bold_path or not regular_path.exists() or not bold_path.exists():
            continue
        try:
            if "PortalSans" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("PortalSans", str(regular_path)))
            if "PortalSansBold" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("PortalSansBold", str(bold_path)))
            return "PortalSans", "PortalSansBold"
        except Exception:
            continue
    return regular, bold

def _metin(value, fallback="-"):
    if value is None:
        return fallback
    value = str(value).strip()
    return value or fallback


def _tarih(value):
    return value.strftime("%d/%m/%Y") if value else "-"


def _izin_gunu(value):
    """Tam günlerde ,00 göstermeden, kesirli günleri koruyarak yazar."""
    if value is None or value == "":
        return "-"
    try:
        return format(Decimal(str(value)).normalize(), "f").replace(".", ",")
    except (InvalidOperation, ValueError, TypeError):
        return _metin(value)


def _telefon_bicimi(value):
    """Cep numarasını PDF'de okunur standart biçimde gösterir."""
    raw = _metin(value, "")
    rakamlar = re.sub(r"\D", "", raw)
    if len(rakamlar) == 11 and rakamlar.startswith("0"):
        rakamlar = rakamlar[1:]
    if len(rakamlar) == 10:
        return f"({rakamlar[:3]}) {rakamlar[3:6]}-{rakamlar[6:8]}-{rakamlar[8:]}"
    return raw or "-"


def _kullanici_adi(user):
    if not user:
        return "Tanımlanmamış"
    return _ad_soyad_bicimi(user.get_full_name().strip() or user.username)


def _kullanici_parafladi_mi(talep, user, rol=None):
    from .imza_durumu import paraf_verildi_mi
    return paraf_verildi_mi(talep, user, rol)


def _izin_aciklamasi(talep):
    izin_turu = talep.izin_turu
    if izin_turu and izin_turu.yillik_izin_mi:
        return (
            f"{_metin(talep.ait_oldugu_yil)} yılı yıllık iznimden "
            f"{_izin_gunu(talep.kullanilan_gun)} gün izin kullanmak istiyorum."
        )
    return f"{_metin(getattr(izin_turu, 'ad', None), 'İzin')} kullanmak istiyorum."


def _birim_satirlari(personel):
    ust_birim = _metin(getattr(personel, "gorev_birimi", ""), "")
    alt_birim = _metin(getattr(personel, "alt_birimi", ""), "")
    return ust_birim, alt_birim


def _belge_gorevi(personel):
    """Belgede işçi/şirket için statü, memur için özlük unvanı gösterilir."""
    if not personel:
        return "-"
    tip = str(getattr(personel, "personel_tipi", "") or "").upper()
    sinif = personel.__class__.__name__.lower()
    if tip == "SIRKET" or "sirket" in sinif:
        return "Şirket Personeli"
    if tip == "ISCI" or "isci" in sinif:
        return "Kadrolu İşçi"
    return _metin(getattr(personel, "unvani", "") or getattr(personel, "unvan", "") or getattr(personel, "gorevi", ""))


def _personel_tipi_etiketi(talep):
    """Belgedeki statüyü okunur başlık biçiminde yazar."""
    return {
        "MEMUR": "Memur",
        "ISCI": "İşçi",
        "SIRKET": "Şirket",
    }.get(str(getattr(talep, "personel_tipi", "") or "").upper(), _metin(talep.get_personel_tipi_display()))


def _p(value, style):
    return Paragraph(escape(_metin(value)), style)


def izin_pdf_olustur(talep, personel, imzacilar, goreve_baslama, islak_imza=False):
    talep, personel, imzacilar = belge_verileri(talep, personel, imzacilar)
    cep_telefonu = getattr(talep, "cep_telefonu", "") or getattr(personel, "cep_tel", "")
    """İzin isteği onayını tek sayfalık, indirilebilir PDF olarak üretir."""
    regular, bold = _fontlar()
    kurum_lacivert = colors.HexColor("#123B59")
    kurum_turkuaz = colors.HexColor("#0A8C91")
    etiket_zemini = colors.HexColor("#EAF3F4")
    yumusak_zemin = colors.HexColor("#F7FAFB")
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=11 * mm,
        bottomMargin=12 * mm,
        title=f"İzin İsteği - {talep.referans_no}",
        author="TİSKİ İzin Portalı",
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle("portal_title", parent=styles["Normal"], fontName=bold, fontSize=13, leading=15, alignment=TA_CENTER, textColor=kurum_lacivert)
    subtitle = ParagraphStyle("portal_subtitle", parent=styles["Normal"], fontName=bold, fontSize=13, leading=15, alignment=TA_CENTER, textColor=kurum_lacivert)
    cell = ParagraphStyle("portal_cell", parent=styles["Normal"], fontName=regular, fontSize=10.5, leading=12.5)
    cell_center = ParagraphStyle("portal_cell_center", parent=cell, alignment=TA_CENTER)
    cell_value = ParagraphStyle("portal_cell_value", parent=cell, alignment=TA_LEFT, leftIndent=8)
    cell_bold = ParagraphStyle("portal_cell_bold", parent=cell, fontName=bold)
    birim_bold = ParagraphStyle("portal_birim_bold", parent=cell_center, fontName=bold, leading=14)
    birim_normal = ParagraphStyle("portal_birim_normal", parent=cell_center, leading=14)
    statu_style = ParagraphStyle("portal_statu", parent=cell, fontSize=9.2, leading=11, alignment=TA_LEFT)
    tiny_left = ParagraphStyle("portal_tiny_left", parent=cell, fontSize=9.8, leading=11.5, alignment=TA_LEFT)
    tiny_center = ParagraphStyle("portal_tiny_center", parent=cell_center, fontSize=10.8, leading=12.5)
    sign_center = ParagraphStyle("portal_sign_center", parent=tiny_center, fontName=bold)
    foot = ParagraphStyle("portal_foot", parent=cell, fontSize=9.4, leading=11)
    story = []

    left_logo = PdfImage(str(BELEDIYE_LOGO), width=19 * mm, height=19 * mm)
    # TİSKİ kaynak görselinin çevresinde boşluk bulunduğu için, iki logo
    # belgede gözle aynı büyüklükte görünsün diye sağ antet biraz büyütülür.
    right_logo = PdfImage(str(TISKI_LOGO), width=23 * mm, height=23 * mm)
    antet = Table(
        [[
            left_logo,
            Paragraph(
                "T.C.<br/><b>TRABZON BÜYÜKŞEHİR BELEDİYESİ</b><br/>"
                "Trabzon İçmesuyu ve Kanalizasyon İdaresi Genel Müdürlüğü",
                title,
            ),
            right_logo,
        ]],
        colWidths=[27 * mm, 126 * mm, 27 * mm],
    )
    antet.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (0, 0), "LEFT"),
        ("ALIGN", (2, 0), (2, 0), "RIGHT"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    story += [antet, Spacer(1, 3 * mm)]

    # Tarih/sayı soldadır; birim adı antette tekrarlanmaz.
    if talep.evrak_no:
        evrak_sayisi = escape(_metin(talep.evrak_no))
        if not islak_imza:
            evrak_sayisi = f"E.{evrak_sayisi}"
        numara_satiri = f"<b>Evrak Tarih ve Sayısı:</b> {_tarih(talep.kurum_kayit_tarihi)} - {evrak_sayisi}"
    else:
        numara_satiri = f"<b>Takip No:</b> {escape(talep.takip_no)}"
    story.append(Paragraph(
        numara_satiri,
        ParagraphStyle("doc_number", parent=cell_bold, alignment=TA_LEFT, fontSize=9.4),
    ))
    story += [Spacer(1, 2 * mm), Paragraph("İZİN İSTEĞİ ONAYI", subtitle), Spacer(1, 3 * mm)]

    ust_birim, alt_birim = _birim_satirlari(personel)
    # Birim adları hücre hizasından bağımsız olarak paragraf seviyesinde ortalanır.
    birim_satirlari = [[Paragraph(escape(ust_birim), birim_bold)]]
    if alt_birim:
        birim_satirlari.append([Paragraph(escape(alt_birim), birim_normal)])
    birim_degeri = Table(birim_satirlari, colWidths=[151 * mm])
    birim_degeri.setStyle(TableStyle([
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 2.8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.8),
        # Daire ve alt birimi ayıran ince çizgi korunur.
        ("LINEBELOW", (0, 0), (0, 0), 0.65, colors.black),
    ]))
    izin_adi = _metin(getattr(talep.izin_turu, "ad", None), "İzin")
    rows = [
        [_p("Birimi", cell_bold), birim_degeri, "", ""],
        [_p("Adı Soyadı", cell_bold), _p(_ad_soyad_bicimi(personel.ad_soyad), cell_value), _p("Ait Olduğu Yıl", cell_bold), _p(talep.ait_oldugu_yil, cell_value)],
        [_p("Görevi", cell_bold), _p(_belge_gorevi(personel), cell_value), _p("İzin Başlangıç Tarihi", cell_bold), _p(_tarih(talep.baslangic_tarihi), cell_value)],
        [_p("T.C. Kimlik No", cell_bold), _p(personel.tc_no, cell_value), _p("İzin Bitiş Tarihi", cell_bold), _p(_tarih(talep.bitis_tarihi), cell_value)],
        [_p("Kurum Sicil No", cell_bold), _p(personel.sicil_no, cell_value), _p("Göreve Başlama Tarihi", cell_bold), _p(_tarih(goreve_baslama), cell_value)],
        [_p("İzin Türü", cell_bold), _p(izin_adi, cell_value), _p("Toplam İzin", cell_bold), _p(_izin_gunu(talep.izin_hakki), cell_value)],
        [_p("Talep Ettiği İzin Süresi", cell_bold), _p(_izin_gunu(talep.kullanilan_gun), cell_value), _p("Kalan İzin Süresi", cell_bold), _p(_izin_gunu(talep.kalan_izin), cell_value)],
    ]
    info_table = Table(rows, colWidths=[29 * mm, 63 * mm, 36 * mm, 52 * mm])
    info_table.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 1.4, colors.black),
        ("INNERGRID", (0, 0), (-1, -1), 0.95, colors.black),
        ("SPAN", (1, 0), (3, 0)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (1, 0), (3, 0), "LEFT"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (1, 0), (3, 0), 0),
        ("RIGHTPADDING", (1, 0), (3, 0), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 4.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4.5),
    ]))
    story.append(info_table)

    note_html = (
        f'<font name="{bold}">Açıklama:</font><br/>'
        + escape(_izin_aciklamasi(talep))
        + "<br/><br/>Gereğini arz ederim."
    )
    note = Table([
        [Paragraph(note_html, cell)],
        [Paragraph(f'<font name="{bold}">Statü:</font> {escape(_personel_tipi_etiketi(talep))}', statu_style)],
    ], colWidths=[180 * mm])
    note.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 1.25, colors.black),
        ("LINEABOVE", (0, 1), (-1, 1), 0.9, colors.black),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(note)

    requester = Table([
        [Paragraph(f'<font name="{bold}">Cep Telefonu:</font> {escape(_telefon_bicimi(cep_telefonu))}', cell)],
    ], colWidths=[180 * mm])
    requester.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 1.25, colors.black),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(requester)

    def imza_paragrafi(baslik, ad_soyad, rol="", durum_satiri=""):
        bosluk = "<br/><br/><br/>" if islak_imza else "<br/><br/>"
        durum = "" if islak_imza or not durum_satiri else f"{escape(durum_satiri)}<br/>"
        return Paragraph(
            f"<b>{escape(baslik)}</b>{bosluk}{durum}{escape(ad_soyad)}<br/>{escape(rol)}",
            sign_center,
        )

    talep_eden = imza_paragrafi(
        "İzin Talebinde Bulunan",
        _ad_soyad_bicimi(personel.ad_soyad),
        _belge_gorevi(personel),
        "e-Paraflıdır",
    )

    # Tek imzacı varsa, talep eden ile onaylayan aynı imza satırında yer alır.
    # Böylece iki imzalı evrakta gereksiz ikinci bir blok oluşmaz.
    if len(imzacilar) == 1:
        rol, user = imzacilar[0]
        onaylayan = imza_paragrafi(
            "Onaylayan",
            _kullanici_adi(user),
            rol,
            "e-İmzalıdır" if talep.durum == "ONAYLANDI" or talep.imza_tamamlanma_tarihi else "İmza Bekliyor",
        )
        iki_imza_table = Table(
            [[talep_eden, onaylayan]],
            colWidths=[90 * mm, 90 * mm],
            rowHeights=[58 * mm],
        )
        iki_imza_table.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 1.4, colors.black),
            ("INNERGRID", (0, 0), (-1, -1), 0.95, colors.black),
            ("BACKGROUND", (0, 0), (-1, -1), colors.white),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ]))
        story.append(iki_imza_table)
    else:
        paraf_kutulari = [talep_eden]
        if len(imzacilar) > 1:
            rol, user = imzacilar[0]
            paraf_kutulari.append(imza_paragrafi(
                "Teklif Eden",
                _kullanici_adi(user),
                rol,
                "e-Paraflıdır" if talep.durum in {"IMZA_SURECINDE", "ONAYLANDI"} or talep.imza_sirasi > 0 else "Paraf Bekliyor",
            ))
        paraf_table = Table([paraf_kutulari], colWidths=[180 * mm / len(paraf_kutulari)] * len(paraf_kutulari), rowHeights=[52 * mm])
        paraf_table.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 1.4, colors.black),
            ("INNERGRID", (0, 0), (-1, -1), 0.95, colors.black),
            ("BACKGROUND", (0, 0), (-1, -1), colors.white),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ]))
        story.append(paraf_table)

        if imzacilar:
            rol, user = imzacilar[-1]
            approval = imza_paragrafi(
                "Onaylayan",
                _kullanici_adi(user),
                rol,
                "e-İmzalıdır" if talep.durum == "ONAYLANDI" or talep.imza_tamamlanma_tarihi else "İmza Bekliyor",
            )
        else:
            approval = Paragraph("<b>Onaylayan</b><br/>Birimi için imzacı tanımlanmamış.", sign_center)
        approval_table = Table([[approval]], colWidths=[180 * mm], rowHeights=[48 * mm])
        approval_table.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 1.4, colors.black),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ]))
        story.append(approval_table)
    story += [Spacer(1, 2 * mm), Paragraph(
        f"<b>Kontrol Eden</b><br/><b>Personel İzin Kontrolü:</b> {escape(_kullanici_adi(talep.kontrol_eden))}", foot
    )]
    doc.build(story)
    return buffer.getvalue()


def izin_excel_olustur(talep, personel, imzacilar, goreve_baslama):
    talep, personel, imzacilar = belge_verileri(talep, personel, imzacilar)
    """Aynı evrak verilerini düzenlenebilir Excel dosyası olarak üretir."""
    cep_telefonu = getattr(talep, "cep_telefonu", "") or getattr(personel, "cep_tel", "")
    wb = Workbook()
    ws = wb.active
    ws.title = "İzin İsteği"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.35
    ws.page_margins.right = 0.35
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35

    for col, width in {"A": 19, "B": 31, "C": 23, "D": 23}.items():
        ws.column_dimensions[col].width = width
    for row in range(1, 34):
        ws.row_dimensions[row].height = 21
    thin = Side(style="thin", color="111111")
    medium = Side(style="medium", color="111111")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    outer_border = Border(left=medium, right=medium, top=medium, bottom=medium)
    no_border = Border()
    label_fill = PatternFill("solid", fgColor="EAF3F5")
    heading = Font(name="Arial", size=12, bold=True)
    heading_big = Font(name="Arial", size=13, bold=True)
    body = Font(name="Arial", size=10)
    body_bold = Font(name="Arial", size=10, bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def style_range(cell_range, *, cell_border=border, fill=None, font=None, alignment=None):
        for row_cells in ws[cell_range]:
            for cell in row_cells:
                cell.border = cell_border
                if fill:
                    cell.fill = fill
                if font:
                    cell.font = font
                if alignment:
                    cell.alignment = alignment

    def merge_write(cell_range, value, *, font=body, alignment=left, fill=None, cell_border=border):
        ws.merge_cells(cell_range)
        cell = ws[cell_range.split(":")[0]]
        cell.value = value
        style_range(cell_range, cell_border=cell_border, fill=fill, font=font, alignment=alignment)
        return cell

    def write_cell(cell_ref, value, *, font=body, alignment=left, fill=None, cell_border=border):
        cell = ws[cell_ref]
        cell.value = value
        cell.font = font
        cell.alignment = alignment
        cell.border = cell_border
        if fill:
            cell.fill = fill
        return cell

    def add_logo(path, anchor, size):
        try:
            image = ExcelImage(str(path))
            image.width, image.height = size, size
            ws.add_image(image, anchor)
        except Exception:
            pass

    add_logo(BELEDIYE_LOGO, "A1", 58)
    add_logo(TISKI_LOGO, "D1", 62)
    merge_write(
        "B1:C1",
        "T.C.\nTRABZON BÜYÜKŞEHİR BELEDİYESİ\nTrabzon İçmesuyu ve Kanalizasyon İdaresi Genel Müdürlüğü",
        font=heading_big,
        alignment=center,
        cell_border=no_border,
    )
    ws["A1"].border = no_border
    ws["D1"].border = no_border
    ws["A1"].alignment = center
    ws["D1"].alignment = center
    ws.row_dimensions[1].height = 63
    numara_satiri = f"Evrak Tarih ve Sayısı: {_tarih(talep.kurum_kayit_tarihi)} - E.{talep.evrak_no}" if talep.evrak_no else f"Takip No: {talep.takip_no}"
    merge_write("A3:D3", numara_satiri, font=body_bold, alignment=left, cell_border=no_border)
    merge_write("A4:D4", "İZİN İSTEĞİ ONAYI", font=heading, alignment=center, cell_border=no_border)
    ws.row_dimensions[4].height = 26

    ust_birim, alt_birim = _birim_satirlari(personel)
    merge_write("A6:A7", "Birimi", font=body_bold, alignment=center, fill=label_fill)
    merge_write("B6:D7", "\n".join(part for part in (ust_birim, alt_birim) if part), font=body_bold, alignment=center)
    values = [
        ("Adı Soyadı", _ad_soyad_bicimi(personel.ad_soyad), "Ait Olduğu Yıl", _metin(talep.ait_oldugu_yil)),
        ("Görevi", _belge_gorevi(personel), "İzin Başlangıç Tarihi", _tarih(talep.baslangic_tarihi)),
        ("T.C. Kimlik No", personel.tc_no, "İzin Bitiş Tarihi", _tarih(talep.bitis_tarihi)),
        ("Kurum Sicil No", personel.sicil_no, "Göreve Başlama Tarihi", _tarih(goreve_baslama)),
        ("İzin Türü", _metin(getattr(talep.izin_turu, "ad", None)), "Toplam İzin", _izin_gunu(talep.izin_hakki)),
        ("Talep Ettiği İzin Süresi", _izin_gunu(talep.kullanilan_gun), "Kalan İzin Süresi", _izin_gunu(talep.kalan_izin)),
    ]
    for row_idx, value in enumerate(values, 8):
        for col_idx, text in enumerate(value, 1):
            write_cell(
                f"{get_column_letter(col_idx)}{row_idx}",
                text,
                font=body_bold if col_idx in {1, 3} else body,
                alignment=center if col_idx in {1, 3} else left,
                fill=label_fill if col_idx in {1, 3} else None,
            )

    note_row = 14
    merge_write(
        f"A{note_row}:D{note_row + 3}",
        (
            "Açıklama:\n"
            f"{_izin_aciklamasi(talep)}\n\n"
            "Gereğini arz ederim.\n"
            f"Statü: {_personel_tipi_etiketi(talep)}        Kalan İzin Süresi: {_izin_gunu(talep.kalan_izin)}"
        ),
        font=body,
        alignment=Alignment(horizontal="left", vertical="top", wrap_text=True),
    )
    for row_idx in range(note_row, note_row + 4):
        ws.row_dimensions[row_idx].height = 24

    phone_row = 18
    merge_write(f"A{phone_row}:D{phone_row}", f"Cep Telefonu: {_telefon_bicimi(cep_telefonu)}", font=body_bold, alignment=left)

    sign_start = 19
    sign_end = 24
    talep_eden = (
        "İzin Talebinde Bulunan\n\n"
        f"{'e-Paraflıdır' if _kullanici_parafladi_mi(talep, talep.kullanici) else 'Paraf Bekliyor'}\n"
        f"{_ad_soyad_bicimi(personel.ad_soyad)}\n"
        f"{_belge_gorevi(personel)}"
    )
    if len(imzacilar) == 1:
        rol, user = imzacilar[0]
        onaylayan = (
            "Onaylayan\n\n"
            f"{'e-İmzalıdır' if talep.durum == 'ONAYLANDI' or talep.imza_tamamlanma_tarihi else 'İmza Bekliyor'}\n"
            f"{_kullanici_adi(user)}\n{rol}"
        )
        merge_write(f"A{sign_start}:B{sign_end}", talep_eden, font=body_bold, alignment=center, cell_border=outer_border)
        merge_write(f"C{sign_start}:D{sign_end}", onaylayan, font=body_bold, alignment=center, cell_border=outer_border)
        row = sign_end + 1
    else:
        merge_write(f"A{sign_start}:B{sign_end}", talep_eden, font=body_bold, alignment=center, cell_border=outer_border)
        if len(imzacilar) > 1:
            rol, user = imzacilar[0]
            teklif_eden = (
                "Teklif Eden\n\n"
                f"{'e-Paraflıdır' if _kullanici_parafladi_mi(talep, user, rol) else 'Paraf Bekliyor'}\n"
                f"{_kullanici_adi(user)}\n{rol}"
            )
            merge_write(f"C{sign_start}:D{sign_end}", teklif_eden, font=body_bold, alignment=center, cell_border=outer_border)
        else:
            merge_write(f"C{sign_start}:D{sign_end}", "", font=body_bold, alignment=center, cell_border=outer_border)
        row = sign_end + 1
        rol, user = imzacilar[-1] if imzacilar else ("", None)
        approval = (
            "Onaylayan\n\n"
            f"{'e-İmzalıdır' if talep.durum == 'ONAYLANDI' or talep.imza_tamamlanma_tarihi else 'İmza Bekliyor'}\n"
            f"{_kullanici_adi(user)}\n{rol}"
        )
        merge_write(f"A{row}:D{row + 4}", approval, font=body_bold, alignment=center, cell_border=outer_border)
        row += 5
    for row_idx in range(sign_start, row):
        ws.row_dimensions[row_idx].height = 23

    merge_write(
        f"A{row}:D{row + 1}",
        f"Kontrol Eden\nPersonel İzin Kontrolü: {_kullanici_adi(talep.kontrol_eden)}",
        font=Font(name="Arial", size=9),
        alignment=left,
        cell_border=no_border,
    )
    ws.print_area = f"A1:D{row + 1}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
