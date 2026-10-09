from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.apps import apps
from django.db import connection, OperationalError
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

from . import test_yetki_akis as akis_fixture
from . import views
from .document_exports import izin_excel_olustur, izin_pdf_olustur
from .models import EvrakNoSayaci, HakEdilenIzin, IzinDenetimKaydi, IzinTalebi, KullanilanIzin
from .services import resmi_evrak_kaydini_hazirla, sonraki_izin_evrak_no


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class IzinNumaralamaTests(TestCase):
    @classmethod
    def setUpClass(cls):
        tables = set(connection.introspection.table_names())
        with connection.schema_editor() as editor:
            for model in apps.get_app_config('portal').get_models():
                if not model._meta.managed and model._meta.db_table not in tables:
                    editor.create_model(model)
                    tables.add(model._meta.db_table)
        super().setUpClass()

    setUpTestData = classmethod(akis_fixture.IzinYetkiAkisTests.setUpTestData.__func__)
    tarihler = akis_fixture.IzinYetkiAkisTests.tarihler
    browser_values = akis_fixture.IzinYetkiAkisTests.browser_values
    gonder = akis_fixture.IzinYetkiAkisTests.gonder
    evrak = akis_fixture.IzinYetkiAkisTests.evrak
    isle = akis_fixture.IzinYetkiAkisTests.isle
    kontrol_asamasina = akis_fixture.IzinYetkiAkisTests.kontrol_asamasina
    imza_asamasina = akis_fixture.IzinYetkiAkisTests.imza_asamasina

    def setUp(self):
        akis_fixture.IzinYetkiAkisTests.setUp(self)
        self.talep.evrak_no = ''
        self.talep.kurum_kayit_tarihi = None
        self.talep.save()
        self.takip = self.talep.takip_no

    def son_imzaya(self):
        self.imza_asamasina()
        self.isle(self.imzaci)
        self.assertEqual(self.talep.kimin_uzerinde_id, self.son_imzaci.pk)

    def sayaclar(self):
        return list(EvrakNoSayaci.objects.order_by('yil').values_list('yil', 'son_no'))

    def tamamla(self):
        self.son_imzaya()
        return self.isle(self.son_imzaci, 'imzala')

    def test_yeni_taslak_acmak_resmi_sayaci_tuketmez(self):
        self.client.force_login(self.sahip)
        before = self.sayaclar()
        response = self.client.get(reverse('izin_evrak', args=['memur', self.personel.pk]))
        self.assertEqual(response.status_code, 302)
        taslak = IzinTalebi.objects.get(durum='TASLAK', kullanici=self.sahip)
        self.assertEqual(taslak.evrak_no, '')
        self.assertIsNone(taslak.kurum_kayit_tarihi)
        self.assertTrue(taslak.takip_no.startswith('TALEP-'))
        self.assertEqual(self.sayaclar(), before)
        self.assertEqual(IzinDenetimKaydi.objects.filter(talep=taslak).latest('pk').evrak_no, taslak.takip_no)

    def test_yeni_taslak_tekrar_acilinca_ayni_takip_kalir(self):
        self.client.force_login(self.sahip)
        url = reverse('izin_evrak', args=['memur', self.personel.pk])
        self.client.get(url)
        taslak = IzinTalebi.objects.get(durum='TASLAK')
        self.client.get(url)
        self.assertEqual(IzinTalebi.objects.filter(durum='TASLAK').count(), 1)
        self.assertEqual(IzinTalebi.objects.get(pk=taslak.pk).takip_no, taslak.takip_no)
        self.assertEqual(self.sayaclar(), [])

    def test_hatali_yeni_form_sayi_ayirmaz(self):
        self.client.force_login(self.sahip)
        response = self.client.post(reverse('izin_evrak', args=['memur', self.personel.pk]), {'islem': 'onaya_sun'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.sayaclar(), [])
        self.assertEqual(IzinTalebi.objects.count(), 1)

    def test_yeni_form_dogrudan_post_ile_takip_alir_resmi_sayi_almaz(self):
        self.client.force_login(self.sahip)
        form = self.client.get(reverse('izin_talebi_duzenle', args=[self.talep.pk])).context['form']
        data = self.browser_values(form)
        data['islem'] = 'onaya_sun'
        response = self.client.post(reverse('izin_evrak', args=['memur', self.personel.pk]), data)
        self.assertEqual(response.status_code, 302)
        yeni = IzinTalebi.objects.exclude(pk=self.talep.pk).get()
        self.assertEqual(yeni.evrak_no, '')
        self.assertIsNone(yeni.kurum_kayit_tarihi)
        self.assertNotEqual(yeni.takip_no, self.takip)
        self.assertEqual(self.sayaclar(), [])

    def test_form_takip_ve_resmi_sayiyi_ayri_gosterir(self):
        self.client.force_login(self.sahip)
        response = self.client.get(reverse('izin_talebi_duzenle', args=[self.talep.pk]))
        self.assertEqual(response.context['form']['takip_no'].value(), self.takip)
        self.assertEqual(response.context['form']['evrak_no'].value(), '')
        self.assertContains(response, 'Son imzada verilecek')

    def test_son_imzaya_kadar_sayi_ve_resmi_tarih_yok(self):
        self.son_imzaya()
        self.assertEqual(self.evrak().evrak_no, '')
        self.assertIsNone(self.talep.kurum_kayit_tarihi)
        self.assertEqual(self.talep.takip_no, self.takip)
        self.assertEqual(self.sayaclar(), [])
        self.assertFalse(KullanilanIzin.objects.exists())

    def test_son_imza_sayi_tarih_ve_ozluk_aktarimini_birlikte_yapar(self):
        self.son_imzaya()
        frozen = self.talep.evrak_verisi.copy()
        today = timezone.localdate()
        self.isle(self.son_imzaci, 'imzala')
        self.assertEqual(self.talep.evrak_no, f'{today.year}-1')
        self.assertEqual(self.talep.kurum_kayit_tarihi, today)
        self.assertEqual(self.talep.takip_no, self.takip)
        self.assertEqual(self.talep.evrak_verisi, frozen)
        self.assertEqual(self.sayaclar(), [(today.year, 1)])
        self.assertEqual(sum(KullanilanIzin.objects.values_list('kullanilan_gun', flat=True)), Decimal('5'))
        self.assertTrue(all(self.talep.evrak_no in item.aciklama for item in KullanilanIzin.objects.all()))

    def test_tekrar_son_imza_yeni_sayi_ayirmaz(self):
        self.tamamla()
        no = self.talep.evrak_no
        before = self.sayaclar()
        kullanim_sayisi = KullanilanIzin.objects.count()
        self.isle(self.son_imzaci, 'imzala')
        self.assertEqual(self.talep.evrak_no, no)
        self.assertEqual(self.sayaclar(), before)
        self.assertEqual(KullanilanIzin.objects.count(), kullanim_sayisi)

    def test_red_yeniden_gonderim_takibi_korur_sayi_tuketmez(self):
        self.imza_asamasina()
        self.isle(self.imzaci, 'reddet', red_nedeni='Duzeltme')
        self.gonder()
        self.assertEqual(self.evrak().takip_no, self.takip)
        self.assertEqual(self.talep.evrak_no, '')
        self.assertIsNone(self.talep.kurum_kayit_tarihi)
        self.assertEqual(self.sayaclar(), [])

    def test_bakiye_yetmezse_sayi_ve_transfer_geri_alinir(self):
        self.son_imzaya()
        HakEdilenIzin.objects.filter(personel_id=self.talep.personel_id).update(hak_edilen_gun=0)
        self.isle(self.son_imzaci, 'imzala')
        self.assertEqual(self.talep.durum, 'IMZA_SURECINDE')
        self.assertEqual(self.talep.evrak_no, '')
        self.assertIsNone(self.talep.kurum_kayit_tarihi)
        self.assertEqual(self.sayaclar(), [])
        self.assertFalse(KullanilanIzin.objects.exists())

    def test_transfer_hatasindan_sonra_tekrar_ilk_resmi_sayiyi_alir(self):
        self.son_imzaya()
        with patch.object(views, 'izni_ozluge_aktar', side_effect=ValueError('test transfer failure')):
            self.isle(self.son_imzaci, 'imzala')
        self.assertEqual(self.sayaclar(), [])
        self.isle(self.son_imzaci, 'imzala')
        self.assertEqual(self.talep.evrak_no, f'{timezone.localdate().year}-1')
        self.assertEqual(sum(KullanilanIzin.objects.values_list('kullanilan_gun', flat=True)), Decimal('5'))

    def test_gecmis_hatasinda_resmi_sayac_da_geri_alinir(self):
        self.son_imzaya()
        self.client.force_login(self.son_imzaci)
        with patch.object(views, '_islem_gecmisi_ekle', side_effect=RuntimeError('test history failure')):
            with self.assertRaises(RuntimeError):
                self.client.post(reverse('evrak_paraf', args=[self.talep.pk]), {'islem': 'imzala'})
        self.assertEqual(self.evrak().evrak_no, '')
        self.assertEqual(self.talep.durum, 'IMZA_SURECINDE')
        self.assertEqual(self.sayaclar(), [])
        self.assertFalse(KullanilanIzin.objects.exists())

    def test_sqlite_kilit_hatasinda_resmi_sayi_da_geri_alinir(self):
        self.son_imzaya()
        self.client.force_login(self.son_imzaci)
        with patch.object(views, '_islem_gecmisi_ekle', side_effect=OperationalError('database is locked')):
            response = self.client.post(reverse('evrak_paraf', args=[self.talep.pk]), {'islem': 'imzala'}, HTTP_ACCEPT='application/json')
        self.assertEqual(response.status_code, 409)
        self.assertEqual(self.evrak().evrak_no, '')
        self.assertEqual(self.sayaclar(), [])
        self.assertFalse(KullanilanIzin.objects.exists())

    def test_yil_degisince_son_imza_yilinin_sayisini_alir(self):
        created = timezone.make_aware(datetime(2026, 12, 30, 10))
        IzinTalebi.objects.filter(pk=self.talep.pk).update(olusturma_zamani=created)
        self.talep.refresh_from_db()
        takip = self.talep.takip_no
        self.son_imzaya()
        with patch('portal.services.timezone.localdate', return_value=date(2027, 1, 2)):
            self.isle(self.son_imzaci, 'imzala')
        self.assertEqual(self.talep.evrak_no, '2027-1')
        self.assertEqual(self.talep.kurum_kayit_tarihi, date(2027, 1, 2))
        self.assertEqual(self.talep.takip_no, takip)
        self.assertEqual(self.sayaclar(), [(2027, 1)])

    def test_mevcut_ayrilmis_sayi_ve_tarih_korunur(self):
        self.talep.evrak_no = '2026-47'
        self.talep.kurum_kayit_tarihi = date(2026, 8, 1)
        self.talep.save()
        EvrakNoSayaci.objects.create(yil=2026, son_no=100)
        self.tamamla()
        self.assertEqual(self.talep.evrak_no, '2026-47')
        self.assertEqual(self.talep.kurum_kayit_tarihi, date(2026, 8, 1))
        self.assertEqual(self.sayaclar(), [(2026, 100)])

    def test_mevcut_sayacin_gerisindeki_sayilar_yeniden_kullanilmaz(self):
        year = timezone.localdate().year
        EvrakNoSayaci.objects.create(yil=year, son_no=100)
        self.tamamla()
        self.assertEqual(self.talep.evrak_no, f'{year}-101')

    def test_sayac_yokken_eski_numaralarin_ustunden_devam_eder(self):
        self.talep.evrak_no = '2026-47'
        self.talep.save()
        self.assertEqual(sonraki_izin_evrak_no(2026), '2026-48')

    def test_silinen_taslak_resmi_sayi_tuketmez(self):
        self.talep.durum = 'TASLAK'
        self.talep.save()
        self.client.force_login(self.sahip)
        response = self.client.post(reverse('izin_talebi_duzenle', args=[self.talep.pk]), {'islem': 'sil'})
        self.assertEqual(response.status_code, 302)
        self.assertFalse(IzinTalebi.objects.filter(pk=self.talep.pk).exists())
        self.assertEqual(self.sayaclar(), [])

    def test_imza_oncesi_pdf_excel_takip_numarasi_gosterir(self):
        self.son_imzaya()
        signers = views.talep_imzacilari(self.talep)
        pdf = izin_pdf_olustur(self.talep, self.personel, signers, self.talep.goreve_baslama_tarihi)
        self.assertTrue(pdf.startswith(b'%PDF'))
        if PdfReader:
            text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(pdf)).pages)
            self.assertIn(self.takip, text)
            self.assertIn('Takip No:', text)
            self.assertNotIn('E.TALEP', text)
        book = load_workbook(BytesIO(izin_excel_olustur(self.talep, self.personel, signers, self.talep.goreve_baslama_tarihi)))
        self.assertEqual(book.active['A3'].value, f'Takip No: {self.takip}')

    def test_imza_sonrasi_pdf_excel_resmi_numara_gosterir(self):
        self.tamamla()
        signers = views.talep_imzacilari(self.talep)
        pdf = izin_pdf_olustur(self.talep, self.personel, signers, self.talep.goreve_baslama_tarihi)
        self.assertTrue(pdf.startswith(b'%PDF'))
        if PdfReader:
            text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(pdf)).pages)
            self.assertIn(f'E.{self.talep.evrak_no}', text)
        book = load_workbook(BytesIO(izin_excel_olustur(self.talep, self.personel, signers, self.talep.goreve_baslama_tarihi)))
        self.assertIn(f'E.{self.talep.evrak_no}', book.active['A3'].value)

    def test_islak_imza_ciktisi_resmi_sayi_vermez(self):
        self.client.force_login(self.sahip)
        data = self.browser_values(self.client.get(reverse('izin_talebi_duzenle', args=[self.talep.pk])).context['form'])
        response = self.client.post(reverse('izin_talebi_islak_cikti', args=[self.talep.pk]), data)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertEqual(self.evrak().evrak_no, '')
        self.assertEqual(self.sayaclar(), [])

    def test_takip_numarasi_resmi_sayi_verildikten_sonra_da_aranir(self):
        self.tamamla()
        self.client.force_login(self.sahip)
        for params in ({'evrak_no': self.takip}, {'q': self.takip}, {'evrak_no': self.talep.evrak_no}):
            response = self.client.get(reverse('izin_ara'), params)
            self.assertEqual(response.status_code, 200)
            self.assertEqual([row['talep'].pk for row in response.context['rows']], [self.talep.pk])

    def test_resmi_tarihi_bos_talep_varsayilan_tarih_filtresinde_kaybolmaz(self):
        self.client.force_login(self.sahip)
        response = self.client.get(reverse('izin_ara'), {'evrak_no': self.takip})
        self.assertEqual([row['talep'].pk for row in response.context['rows']], [self.talep.pk])

    def test_gecersiz_takip_numarasi_500_uretmez(self):
        self.client.force_login(self.sahip)
        for value in ('TALEP-0000-1', 'TALEP-2026-' + '9' * 100, 'TALEP-2026-' + '9' * 5000, 'TALEP-2026-x'):
            self.assertEqual(self.client.get(reverse('izin_ara'), {'evrak_no': value}).status_code, 200)

    def test_takip_numarasi_siralamasi_sayfalar_arasinda_tutarlidir(self):
        IzinTalebi.objects.bulk_create([
            IzinTalebi(kullanici=self.sahip, personel_id=self.talep.personel_id, personel_tipi='MEMUR', durum='REDDEDILDI')
            for _ in range(15)
        ])
        expected = sorted(IzinTalebi.objects.values_list('pk', flat=True), key=lambda pk: IzinTalebi.objects.get(pk=pk).takip_no)
        self.client.force_login(self.sahip)
        for direction in ('asc', 'desc'):
            rows = []
            for page in (1, 2):
                response = self.client.get(reverse('izin_ara'), {'sort': 'evrak_no', 'dir': direction, 'page': page, 'per_page': 10})
                self.assertEqual(response.status_code, 200)
                rows.extend(row['talep'].pk for row in response.context['rows'])
            self.assertEqual(rows, expected if direction == 'asc' else list(reversed(expected)))

    def test_son_imza_disinda_resmi_sayi_hazirlanamaz(self):
        with self.assertRaises(ValueError):
            resmi_evrak_kaydini_hazirla(self.talep)
        self.assertEqual(self.sayaclar(), [])
